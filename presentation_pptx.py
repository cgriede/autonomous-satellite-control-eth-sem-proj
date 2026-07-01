"""Read, validate, edit, and build project presentation decks from a JSON manifest."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

MANIFEST_VERSION = 1

ALLOWED_SECTIONS = frozenset({"main", "admin", "backup"})
ALLOWED_STATUS = frozenset({"draft", "ready", "placeholder"})

DEFAULT_MANIFEST_DIR = (
    Path(__file__).resolve().parents[2] / "docs" / "presentation" / "final" / "with-cursor"
)
DEFAULT_MANIFEST = DEFAULT_MANIFEST_DIR / "slides_manifest.json"
DEFAULT_PPTX = DEFAULT_MANIFEST_DIR / "work_review_with_cursor.pptx"

REQUIRED_SLIDE_KEYS = frozenset({"id", "section", "title", "bullets", "status"})


def _load_manifest(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, dict):
        raise ValueError("Manifest root must be a JSON object")
    if "slides" not in data or not isinstance(data["slides"], list):
        raise ValueError("Manifest must contain a 'slides' array")
    return data


def _save_manifest(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)
        fh.write("\n")


def _validate_slide(slide: dict[str, Any], *, row_number: int) -> None:
    missing = REQUIRED_SLIDE_KEYS - slide.keys()
    if missing:
        raise ValueError(f"Slide row {row_number}: missing keys {sorted(missing)}")
    if not slide["id"]:
        raise ValueError(f"Slide row {row_number}: id is required")
    section = slide["section"]
    if section not in ALLOWED_SECTIONS:
        raise ValueError(f"Slide {slide['id']}: invalid section {section!r}")
    status = slide["status"]
    if status not in ALLOWED_STATUS:
        raise ValueError(f"Slide {slide['id']}: invalid status {status!r}")
    bullets = slide["bullets"]
    if not isinstance(bullets, list) or not all(isinstance(b, str) for b in bullets):
        raise ValueError(f"Slide {slide['id']}: bullets must be a list of strings")


def assert_manifest_healthy(path: Path | str) -> list[dict[str, Any]]:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Manifest not found: {path}")
    data = _load_manifest(path)
    slides = data["slides"]
    if not slides:
        raise ValueError(f"Manifest has no slides: {path}")
    seen: set[str] = set()
    for row_number, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            raise ValueError(f"Slide row {row_number}: must be an object")
        _validate_slide(slide, row_number=row_number)
        slide_id = slide["id"]
        if slide_id in seen:
            raise ValueError(f"Duplicate slide id {slide_id!r}")
        seen.add(slide_id)
    return slides


def read_slides(path: Path | str) -> list[dict[str, Any]]:
    return assert_manifest_healthy(path)


def set_slide(
    path: Path | str,
    slide_id: str,
    *,
    title: str | None = None,
    bullets: list[str] | None = None,
    notes: str | None = None,
    status: str | None = None,
) -> dict[str, Any]:
    path = Path(path)
    data = _load_manifest(path)
    for slide in data["slides"]:
        if slide["id"] != slide_id:
            continue
        if title is not None:
            slide["title"] = title
        if bullets is not None:
            slide["bullets"] = bullets
        if notes is not None:
            slide["notes"] = notes
        if status is not None:
            slide["status"] = status
        _save_manifest(path, data)
        assert_manifest_healthy(path)
        return slide
    raise KeyError(f"Slide id not found: {slide_id!r}")


def append_bullets(path: Path | str, slide_id: str, new_bullets: list[str]) -> dict[str, Any]:
    slides = read_slides(path)
    slide = next(s for s in slides if s["id"] == slide_id)
    merged = list(slide["bullets"]) + list(new_bullets)
    return set_slide(path, slide_id, bullets=merged)


def _add_bullet_slide(prs: Presentation, title: str, bullets: list[str], notes: str = "") -> None:
    layout = prs.slide_layouts[1]  # title + content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    if not bullets:
        body.text = "[content TBD]"
        return
    body.text = bullets[0]
    for bullet in bullets[1:]:
        p = body.add_paragraph()
        p.text = bullet
        p.level = 0
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _add_image_slide(
    prs: Presentation,
    title: str,
    image_path: Path,
    *,
    bullets: list[str] | None = None,
    notes: str = "",
) -> None:
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.6)).text_frame.text = title
    slide.shapes.add_picture(str(image_path), Inches(0.35), Inches(0.85), width=Inches(12.6))
    if bullets:
        box = slide.shapes.add_textbox(Inches(0.4), Inches(6.85), Inches(12.5), Inches(0.55))
        tf = box.text_frame
        tf.text = bullets[0]
        for bullet in bullets[1:]:
            p = tf.add_paragraph()
            p.text = bullet
            p.level = 0
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def build_pptx(manifest_path: Path | str, output_path: Path | str) -> Path:
    manifest_path = Path(manifest_path)
    output_path = Path(output_path)
    data = _load_manifest(manifest_path)
    slides = assert_manifest_healthy(manifest_path)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    deck_title = data.get("deck_title", "Presentation")
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = deck_title
    subtitle = title_slide.placeholders[1]
    subtitle.text = data.get("deck_subtitle", "")

    current_section: str | None = None
    for slide in slides:
        section = slide["section"]
        if section != current_section:
            current_section = section

        notes = slide.get("notes", "")
        status = slide.get("status", "draft")
        if status == "placeholder":
            bullets = list(slide["bullets"])
            if bullets and bullets[-1] != "[YOUR CONTENT]":
                bullets.append("[YOUR CONTENT]")
        else:
            bullets = slide["bullets"]

        image_rel = slide.get("image")
        if image_rel:
            image_path = Path(image_rel)
            if not image_path.is_absolute():
                image_path = manifest_path.parent / image_path
            if not image_path.is_file():
                raise FileNotFoundError(
                    f"Slide {slide['id']}: image not found: {image_path} "
                    f"(run: python backend/scripts/presentation_diagrams.py)"
                )
            _add_image_slide(
                prs,
                f"{slide['id']}: {slide['title']}",
                image_path,
                bullets=bullets if bullets else None,
                notes=notes,
            )
            continue

        _add_bullet_slide(prs, f"{slide['id']}: {slide['title']}", bullets, notes=notes)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path


def _print_slides(slides: list[dict[str, Any]]) -> None:
    sections: dict[str, int] = {}
    for slide in slides:
        sections[slide["section"]] = sections.get(slide["section"], 0) + 1
    print(
        "slides={n} main={main} admin={admin} backup={backup}".format(
            n=len(slides),
            main=sections.get("main", 0),
            admin=sections.get("admin", 0),
            backup=sections.get("backup", 0),
        )
    )
    for slide in slides:
        bullet_preview = slide["bullets"][0][:60] + "…" if slide["bullets"] else "(empty)"
        print(
            f"{slide['id']:4}  {slide['section']:6}  {slide['status']:11}  "
            f"{slide['title'][:50]:50}  | {bullet_preview}"
        )


def init_default_manifest(path: Path | str, *, overwrite: bool = False) -> Path:
    path = Path(path)
    if path.exists() and not overwrite:
        raise FileExistsError(f"Manifest already exists: {path} (use --force)")
    from scripts.presentation_seed_final_review import DEFAULT_FINAL_REVIEW_MANIFEST

    _save_manifest(path, DEFAULT_FINAL_REVIEW_MANIFEST)
    assert_manifest_healthy(path)
    return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Presentation manifest + PPTX utilities.")
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("check")
    sub.add_parser("list")
    sub.add_parser("build")

    diag_p = sub.add_parser("diagrams", help="Render flowchart PNGs from presentation_diagrams.py")

    init_p = sub.add_parser("init")
    init_p.add_argument(
        "--force",
        action="store_true",
        help="Overwrite existing manifest",
    )

    set_p = sub.add_parser("set", help="Update one slide in the manifest")
    set_p.add_argument("--id", required=True, help="Slide id (e.g. M12)")
    set_p.add_argument("--title")
    set_p.add_argument("--status", choices=sorted(ALLOWED_STATUS))
    set_p.add_argument("--notes")
    set_p.add_argument("--bullets", help="Pipe-separated bullet lines")

    append_p = sub.add_parser("append", help="Append bullets to a slide")
    append_p.add_argument("--id", required=True)
    append_p.add_argument("--bullets", required=True, help="Pipe-separated bullet lines")

    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--output", type=Path, default=DEFAULT_PPTX)

    args = parser.parse_args(argv)
    manifest = args.manifest

    if args.command == "init":
        init_default_manifest(manifest, overwrite=getattr(args, "force", False))
        print(f"Created {manifest}")
        return 0

    if args.command == "check":
        slides = assert_manifest_healthy(manifest)
        print(f"OK: {manifest} ({len(slides)} slides)")
        return 0

    if args.command == "list":
        _print_slides(read_slides(manifest))
        return 0

    if args.command == "build":
        out = build_pptx(manifest, args.output)
        print(f"Wrote {out}")
        return 0

    if args.command == "diagrams":
        from scripts.presentation_diagrams import render_all

        diagram_dir = manifest.parent / "diagrams"
        ml, sim = render_all(diagram_dir)
        print(f"Wrote {ml}")
        print(f"Wrote {sim}")
        return 0

    if args.command == "set":
        bullets = None
        if args.bullets is not None:
            bullets = [b.strip() for b in args.bullets.split("|") if b.strip()]
        set_slide(
            manifest,
            args.id,
            title=args.title,
            bullets=bullets,
            notes=args.notes,
            status=args.status,
        )
        print(f"Updated {args.id}")
        return 0

    if args.command == "append":
        bullets = [b.strip() for b in args.bullets.split("|") if b.strip()]
        append_bullets(manifest, args.id, bullets)
        print(f"Appended to {args.id}")
        return 0

    raise AssertionError(f"Unhandled command: {args.command}")


if __name__ == "__main__":
    sys.exit(main())
