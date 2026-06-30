"""Read, validate, edit, and build project presentation decks from a JSON manifest.

Agent tool — canonical location: .cursor/tools/presentation/
See .cursor/skills/build-tool/SKILL.md
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

_TOOL_DIR = Path(__file__).resolve().parent
if str(_TOOL_DIR) not in sys.path:
    sys.path.insert(0, str(_TOOL_DIR))

_REPO_ROOT = Path(__file__).resolve().parents[3]

MANIFEST_VERSION = 1

ALLOWED_SECTIONS = frozenset({"main", "admin", "backup"})
ALLOWED_STATUS = frozenset({"draft", "ready", "placeholder"})

DEFAULT_MANIFEST_DIR = _REPO_ROOT / "docs" / "presentation" / "final" / "with-cursor"
DEFAULT_MANIFEST = DEFAULT_MANIFEST_DIR / "slides_manifest.json"
DEFAULT_PPTX = DEFAULT_MANIFEST_DIR / "work_review_with_cursor.pptx"

REQUIRED_SLIDE_KEYS = frozenset({"id", "section", "title", "bullets", "status"})

# --- visual theme (dark "mission-control" deck) ------------------------------

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

C_BG = RGBColor(0x0B, 0x12, 0x21)        # deep space navy
C_PANEL = RGBColor(0x16, 0x22, 0x3B)     # raised card
C_INK = RGBColor(0xEA, 0xF0, 0xFB)       # primary text
C_MUTED = RGBColor(0x93, 0xA4, 0xC3)     # secondary text
C_CYAN = RGBColor(0x38, 0xBD, 0xF8)      # primary accent
C_GREEN = RGBColor(0x34, 0xD3, 0x99)
C_AMBER = RGBColor(0xFB, 0xBF, 0x24)
C_EDGE = RGBColor(0x2C, 0x3E, 0x63)

FONT = "Segoe UI"

# accent colour per deck section
SECTION_ACCENT = {"main": C_CYAN, "admin": C_AMBER, "backup": C_GREEN}


def _load_manifest(path: Path) -> dict[str, Any]:
    with path.open(encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, dict):
        raise ValueError("Manifest root must be a JSON object")
    if "slides" not in data or not isinstance(data["slides"], list):
        raise ValueError("Manifest must contain a 'slides' array")
    return data


def _save_manifest(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)
        fh.write("\n")


def _validate_slide(slide: dict[str, Any], *, row_number: int) -> None:
    missing = REQUIRED_SLIDE_KEYS - slide.keys()
    if missing:
        raise ValueError(f"Slide row {row_number}: missing keys {sorted(missing)}")
    if not slide["id"]:
        raise ValueError(f"Slide row {row_number}: id is required")
    section = slide["section"]
    if section not in ALLOWED_SECTIONS:
        raise ValueError(f"Slide {slide['id']}: invalid section {section!r}")
    status = slide["status"]
    if status not in ALLOWED_STATUS:
        raise ValueError(f"Slide {slide['id']}: invalid status {status!r}")
    bullets = slide["bullets"]
    if not isinstance(bullets, list) or not all(isinstance(b, str) for b in bullets):
        raise ValueError(f"Slide {slide['id']}: bullets must be a list of strings")


def assert_manifest_healthy(path: Path | str) -> list[dict[str, Any]]:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Manifest not found: {path}")
    data = _load_manifest(path)
    slides = data["slides"]
    if not slides:
        raise ValueError(f"Manifest has no slides: {path}")
    seen: set[str] = set()
    for row_number, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            raise ValueError(f"Slide row {row_number}: must be an object")
        _validate_slide(slide, row_number=row_number)
        slide_id = slide["id"]
        if slide_id in seen:
            raise ValueError(f"Duplicate slide id {slide_id!r}")
        seen.add(slide_id)
    return slides


def read_slides(path: Path | str) -> list[dict[str, Any]]:
    return assert_manifest_healthy(path)


def set_slide(
    path: Path | str,
    slide_id: str,
    *,
    title: str | None = None,
    bullets: list[str] | None = None,
    notes: str | None = None,
    status: str | None = None,
) -> dict[str, Any]:
    path = Path(path)
    data = _load_manifest(path)
    for slide in data["slides"]:
        if slide["id"] != slide_id:
            continue
        if title is not None:
            slide["title"] = title
        if bullets is not None:
            slide["bullets"] = bullets
        if notes is not None:
            slide["notes"] = notes
        if status is not None:
            slide["status"] = status
        _save_manifest(path, data)
        assert_manifest_healthy(path)
        return slide
    raise KeyError(f"Slide id not found: {slide_id!r}")


def append_bullets(path: Path | str, slide_id: str, new_bullets: list[str]) -> dict[str, Any]:
    slides = read_slides(path)
    slide = next(s for s in slides if s["id"] == slide_id)
    merged = list(slide["bullets"]) + list(new_bullets)
    return set_slide(path, slide_id, bullets=merged)


def _png_size(path: Path) -> tuple[int, int]:
    """Return (width, height) px for an image, preferring Pillow."""
    try:
        from PIL import Image

        with Image.open(path) as im:
            return im.size
    except Exception:
        # Minimal PNG IHDR fallback (width/height at bytes 16-24).
        with path.open("rb") as fh:
            head = fh.read(24)
        if len(head) >= 24 and head[:8] == b"\x89PNG\r\n\x1a\n":
            w = int.from_bytes(head[16:20], "big")
            h = int.from_bytes(head[20:24], "big")
            return w, h
        return (1600, 900)


def _blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return slide


def _rect(slide, x, y, w, h, color, *, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def _text(slide, x, y, w, h, runs, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          space_after=6, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph is a list of (text, size, color, bold)."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for text, size, color, bold in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = FONT
    return box


def _footer(slide, deck_title, idx, accent):
    _rect(slide, Inches(0.0), Inches(7.18), SLIDE_W, Pt(2.2), accent)
    _text(slide, Inches(0.45), Inches(7.16), Inches(9.0), Inches(0.3),
          [[(deck_title, 9, C_MUTED, False)]], anchor=MSO_ANCHOR.MIDDLE)
    _text(slide, Inches(11.4), Inches(7.16), Inches(1.5), Inches(0.3),
          [[(f"{idx:02d}", 9, accent, True)]], align=PP_ALIGN.RIGHT,
          anchor=MSO_ANCHOR.MIDDLE)


def _header(slide, slide_id, title, accent):
    # accent tab + slide id
    _rect(slide, Inches(0.45), Inches(0.5), Inches(0.12), Inches(0.62), accent)
    _text(slide, Inches(0.72), Inches(0.42), Inches(11.5), Inches(0.5),
          [[(f"{slide_id}", 12, accent, True), (f"   {title}", 22, C_INK, True)]],
          anchor=MSO_ANCHOR.MIDDLE)
    _rect(slide, Inches(0.72), Inches(1.16), Inches(2.5), Pt(2.4), accent)
    _rect(slide, Inches(3.22), Inches(1.18), Inches(9.65), Pt(1.0), C_EDGE)


def _add_title_slide(prs, deck_title, subtitle, accent):
    slide = _blank_slide(prs)
    _rect(slide, Inches(0.0), Inches(0.0), Inches(0.22), SLIDE_H, accent)
    _rect(slide, Inches(0.9), Inches(2.55), Inches(3.2), Pt(3.5), accent)
    _text(slide, Inches(0.9), Inches(2.75), Inches(11.5), Inches(2.0),
          [[(deck_title, 40, C_INK, True)]])
    _text(slide, Inches(0.92), Inches(4.35), Inches(11.0), Inches(1.0),
          [[(subtitle, 18, C_CYAN, False)]])
    _text(slide, Inches(0.92), Inches(6.6), Inches(11.0), Inches(0.5),
          [[("Cedric Grieder  ·  ETH x Beyond Gravity semester project", 12, C_MUTED, False)]])


def _add_section_slide(prs, slide_id, title, subtitle, accent, idx, deck_title):
    slide = _blank_slide(prs)
    _rect(slide, Inches(0.0), Inches(3.05), SLIDE_W, Inches(1.6), C_PANEL)
    _rect(slide, Inches(0.9), Inches(3.05), Inches(0.14), Inches(1.6), accent)
    _text(slide, Inches(1.25), Inches(2.35), Inches(11.0), Inches(0.6),
          [[(slide_id, 16, accent, True)]])
    _text(slide, Inches(1.25), Inches(2.95), Inches(11.4), Inches(1.0),
          [[(title, 34, C_INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        _text(slide, Inches(1.27), Inches(4.7), Inches(11.0), Inches(0.6),
              [[(subtitle, 15, C_MUTED, False)]])
    _footer(slide, deck_title, idx, accent)


def _add_bullet_slide(prs, slide_id, title, bullets, accent, idx, deck_title, notes=""):
    slide = _blank_slide(prs)
    _header(slide, slide_id, title, accent)
    paras = []
    for b in bullets:
        paras.append([("▸  ", 14, accent, True), (b, 15, C_INK, False)])
    if not paras:
        paras = [[("[content TBD]", 15, C_MUTED, False)]]
    _text(slide, Inches(0.85), Inches(1.55), Inches(11.7), Inches(5.3), paras,
          space_after=14, line_spacing=1.05)
    _footer(slide, deck_title, idx, accent)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _add_image_slide(prs, slide_id, title, image_path, accent, idx, deck_title,
                     *, caption=None, notes=""):
    slide = _blank_slide(prs)
    _header(slide, slide_id, title, accent)

    has_cap = bool(caption)
    area_x, area_y = Inches(0.45), Inches(1.4)
    area_w = Inches(12.45)
    area_h = Inches(5.0 if has_cap else 5.55)

    px_w, px_h = _png_size(image_path)
    aspect = px_w / px_h if px_h else 1.6
    w = area_w
    h = Emu(int(w / aspect))
    if h > area_h:
        h = area_h
        w = Emu(int(h * aspect))
    x = Emu(int(area_x + (area_w - w) / 2))
    y = Emu(int(area_y + (area_h - h) / 2))
    slide.shapes.add_picture(str(image_path), x, y, width=w, height=h)

    if has_cap:
        cap_runs = [[("▸  ", 12, accent, True), (c, 12.5, C_MUTED, False)] for c in caption]
        _text(slide, Inches(0.85), Inches(6.5), Inches(11.7), Inches(0.6), cap_runs,
              space_after=2, line_spacing=1.0)
    _footer(slide, deck_title, idx, accent)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _add_table_slide(prs, slide_id, title, rows, accent, idx, deck_title,
                     *, intro=None, notes="", col_widths=None):
    """Render a native PowerPoint table (first row = header) on the dark theme."""
    slide = _blank_slide(prs)
    _header(slide, slide_id, title, accent)

    top = Inches(1.55)
    if intro:
        _text(slide, Inches(0.85), Inches(1.45), Inches(11.7), Inches(0.5),
              [[(intro, 13, C_MUTED, False)]])
        top = Inches(2.0)

    n_rows = len(rows)
    n_cols = max(len(r) for r in rows) if rows else 1
    table_w = Inches(12.4)
    avail_h = Inches(6.85) - top
    gfx = slide.shapes.add_table(n_rows, n_cols, Inches(0.46), top, table_w, avail_h)
    table = gfx.table
    table.first_row = False
    table.horz_banding = False

    if col_widths and len(col_widths) == n_cols:
        total = sum(col_widths)
        for c, frac in enumerate(col_widths):
            table.columns[c].width = Emu(int(int(table_w) * frac / total))

    for r, row in enumerate(rows):
        is_head = r == 0
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if is_head:
                cell.fill.fore_color.rgb = accent
            else:
                cell.fill.fore_color.rgb = C_PANEL if r % 2 else RGBColor(0x10, 0x1A, 0x30)
            text = row[c] if c < len(row) else ""
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = text
            run.font.name = FONT
            run.font.size = Pt(12 if is_head else 11)
            run.font.bold = is_head or c == 0
            run.font.color.rgb = C_BG if is_head else C_INK

    _footer(slide, deck_title, idx, accent)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _add_video_slide(prs, slide_id, title, video_path, poster_path, bullets,
                     accent, idx, deck_title, *, notes=""):
    """Bullets on the left, an embedded (click-to-play) video on the right."""
    slide = _blank_slide(prs)
    _header(slide, slide_id, title, accent)

    has_text = bool(bullets)
    if has_text:
        paras = [[("▸  ", 14, accent, True), (b, 14, C_INK, False)] for b in bullets]
        _text(slide, Inches(0.85), Inches(1.55), Inches(5.2), Inches(5.2), paras,
              space_after=12, line_spacing=1.05)
        area_x, area_w = Inches(6.35), Inches(6.55)
    else:
        area_x, area_w = Inches(1.2), Inches(10.9)
    area_y, area_h = Inches(1.55), Inches(5.1)

    aspect = 16 / 9
    if poster_path and Path(poster_path).is_file():
        pw, ph = _png_size(Path(poster_path))
        aspect = pw / ph if ph else aspect
    w = area_w
    h = Emu(int(w / aspect))
    if h > area_h:
        h = area_h
        w = Emu(int(h * aspect))
    x = Emu(int(area_x + (area_w - w) / 2))
    y = Emu(int(area_y + (area_h - h) / 2))

    poster = str(poster_path) if (poster_path and Path(poster_path).is_file()) else None
    try:
        slide.shapes.add_movie(
            str(video_path), x, y, w, h,
            poster_frame_image=poster, mime_type="video/mp4",
        )
    except Exception:
        if poster:
            slide.shapes.add_picture(poster, x, y, width=w, height=h)
        _rect(slide, x, Emu(int(y + h)), w, Inches(0.32), C_PANEL)

    _text(slide, x, Emu(int(y + h + Inches(0.02))), w, Inches(0.34),
          [[("▶  click to play  ·  ", 11, accent, True),
            (Path(video_path).name, 11, C_MUTED, False)]],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    _footer(slide, deck_title, idx, accent)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _resolve(manifest_path: Path, rel: str) -> Path:
    p = Path(rel)
    return p if p.is_absolute() else (manifest_path.parent / p)


def build_pptx(manifest_path: Path | str, output_path: Path | str) -> Path:
    manifest_path = Path(manifest_path)
    output_path = Path(output_path)
    data = _load_manifest(manifest_path)
    slides = assert_manifest_healthy(manifest_path)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    deck_title = data.get("deck_title", "Presentation")
    _add_title_slide(prs, deck_title, data.get("deck_subtitle", ""), C_CYAN)

    for idx, slide in enumerate(slides, start=1):
        section = slide["section"]
        accent = SECTION_ACCENT.get(section, C_CYAN)
        title = slide["title"]
        notes = slide.get("notes", "")
        status = slide.get("status", "draft")

        if slide.get("kind") == "section":
            subtitle = slide["bullets"][0] if slide["bullets"] else ""
            _add_section_slide(prs, slide["id"], title, subtitle, accent, idx, deck_title)
            continue

        if status == "placeholder":
            bullets = list(slide["bullets"])
            if bullets and bullets[-1] != "[YOUR CONTENT]":
                bullets.append("[YOUR CONTENT]")
        else:
            bullets = slide["bullets"]

        table_rows = slide.get("table")
        if table_rows:
            _add_table_slide(
                prs, slide["id"], title, table_rows, accent, idx, deck_title,
                intro=(bullets[0] if bullets else None), notes=notes,
                col_widths=slide.get("col_widths"),
            )
            continue

        video_rel = slide.get("video")
        if video_rel:
            video_path = _resolve(manifest_path, video_rel)
            if not video_path.is_file():
                raise FileNotFoundError(
                    f"Slide {slide['id']}: video not found: {video_path}"
                )
            poster_rel = slide.get("poster")
            poster_path = _resolve(manifest_path, poster_rel) if poster_rel else None
            _add_video_slide(
                prs, slide["id"], title, video_path, poster_path, bullets,
                accent, idx, deck_title, notes=notes,
            )
            continue

        image_rel = slide.get("image")
        if image_rel:
            image_path = Path(image_rel)
            if not image_path.is_absolute():
                image_path = manifest_path.parent / image_path
            if not image_path.is_file():
                raise FileNotFoundError(
                    f"Slide {slide['id']}: image not found: {image_path} "
                    f"(run: python .cursor/tools/presentation/presentation_pptx.py diagrams)"
                )
            _add_image_slide(
                prs, slide["id"], title, image_path, accent, idx, deck_title,
                caption=bullets if bullets else None, notes=notes,
            )
            continue

        _add_bullet_slide(prs, slide["id"], title, bullets, accent, idx, deck_title,
                          notes=notes)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path


def _print_slides(slides: list[dict[str, Any]]) -> None:
    sections: dict[str, int] = {}
    for slide in slides:
        sections[slide["section"]] = sections.get(slide["section"], 0) + 1
    print(
        "slides={n} main={main} admin={admin} backup={backup}".format(
            n=len(slides),
            main=sections.get("main", 0),
            admin=sections.get("admin", 0),
            backup=sections.get("backup", 0),
        )
    )
    for slide in slides:
        bullet_preview = slide["bullets"][0][:60] + "…" if slide["bullets"] else "(empty)"
        print(
            f"{slide['id']:4}  {slide['section']:6}  {slide['status']:11}  "
            f"{slide['title'][:50]:50}  | {bullet_preview}"
        )


def init_default_manifest(path: Path | str, *, overwrite: bool = False) -> Path:
    path = Path(path)
    if path.exists() and not overwrite:
        raise FileExistsError(f"Manifest already exists: {path} (use --force)")
    from presentation_seed_final_review import DEFAULT_FINAL_REVIEW_MANIFEST

    _save_manifest(path, DEFAULT_FINAL_REVIEW_MANIFEST)
    assert_manifest_healthy(path)
    return path


def main(argv: list[str] | None = None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # Windows console safety
    except Exception:
        pass
    parser = argparse.ArgumentParser(description="Presentation manifest + PPTX utilities.")
    sub = parser.add_subparsers(dest="command", required=True)

    sub.add_parser("check")
    sub.add_parser("list")
    sub.add_parser("build")

    diag_p = sub.add_parser("diagrams", help="Render flowchart PNGs from presentation_diagrams.py")

    init_p = sub.add_parser("init")
    init_p.add_argument(
        "--force",
        action="store_true",
        help="Overwrite existing manifest",
    )

    set_p = sub.add_parser("set", help="Update one slide in the manifest")
    set_p.add_argument("--id", required=True, help="Slide id (e.g. M12)")
    set_p.add_argument("--title")
    set_p.add_argument("--status", choices=sorted(ALLOWED_STATUS))
    set_p.add_argument("--notes")
    set_p.add_argument("--bullets", help="Pipe-separated bullet lines")

    append_p = sub.add_parser("append", help="Append bullets to a slide")
    append_p.add_argument("--id", required=True)
    append_p.add_argument("--bullets", required=True, help="Pipe-separated bullet lines")

    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--output", type=Path, default=DEFAULT_PPTX)

    args = parser.parse_args(argv)
    manifest = args.manifest

    if args.command == "init":
        init_default_manifest(manifest, overwrite=getattr(args, "force", False))
        print(f"Created {manifest}")
        return 0

    if args.command == "check":
        slides = assert_manifest_healthy(manifest)
        print(f"OK: {manifest} ({len(slides)} slides)")
        return 0

    if args.command == "list":
        _print_slides(read_slides(manifest))
        return 0

    if args.command == "build":
        out = build_pptx(manifest, args.output)
        print(f"Wrote {out}")
        return 0

    if args.command == "diagrams":
        from presentation_diagrams import render_all

        diagram_dir = manifest.parent / "diagrams"
        for path in render_all(diagram_dir):
            print(f"Wrote {path}")
        return 0

    if args.command == "set":
        bullets = None
        if args.bullets is not None:
            bullets = [b.strip() for b in args.bullets.split("|") if b.strip()]
        set_slide(
            manifest,
            args.id,
            title=args.title,
            bullets=bullets,
            notes=args.notes,
            status=args.status,
        )
        print(f"Updated {args.id}")
        return 0

    if args.command == "append":
        bullets = [b.strip() for b in args.bullets.split("|") if b.strip()]
        append_bullets(manifest, args.id, bullets)
        print(f"Appended to {args.id}")
        return 0

    raise AssertionError(f"Unhandled command: {args.command}")


if __name__ == "__main__":
    sys.exit(main())
