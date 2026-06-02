"""Read, validate, and initialize the live project backlog workbook."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

BACKLOG_SHEET = "backlog"

BACKLOG_COLUMNS: tuple[str, ...] = (
    "uid",
    "Sprint",
    "prio",
    "Size",
    "status",
    "dep on",
    "name",
    "notes / blockers",
)

ALLOWED_STATUS = frozenset({"todo", "wip", "done", "blocked", "icebox"})
ALLOWED_PRIO = frozenset({"P0", "P1", "P2", "P3"})
ALLOWED_SIZE = frozenset({"S", "M", "L", "XL"})

DEFAULT_BACKLOG_ROWS: tuple[dict[str, str], ...] = (
    {
        "uid": "BUG-001",
        "Sprint": "S01",
        "prio": "P1",
        "Size": "M",
        "status": "todo",
        "dep on": "",
        "name": "1d strip updated incorrectly",
        "notes / blockers": "Long-standing; strip semantics still wrong in camera observation path.",
    },
    {
        "uid": "BUG-002",
        "Sprint": "S01",
        "prio": "P1",
        "Size": "M",
        "status": "wip",
        "dep on": "FEAT-005",
        "name": "Clouds not visibly moving in simulation",
        "notes / blockers": "Formation generator exists (notebook 02); kinematic motion / render feedback still weak.",
    },
    {
        "uid": "BUG-003",
        "Sprint": "S01",
        "prio": "P0",
        "Size": "S",
        "status": "wip",
        "dep on": "",
        "name": "Integration tests out of sync with SimulationConfig / cloud precompute",
        "notes / blockers": "movement_constraints e2e (controller_mode removed); cloud formation e2e; test_cloud_arc_precompute (2 red).",
    },
    {
        "uid": "STR-001",
        "Sprint": "",
        "prio": "P2",
        "Size": "L",
        "status": "todo",
        "dep on": "",
        "name": "Reduce mission/global constant coupling in training_runtime",
        "notes / blockers": "Pass scenario/config explicitly instead of scraping MISSION globals.",
    },
    {
        "uid": "STR-002",
        "Sprint": "",
        "prio": "P2",
        "Size": "M",
        "status": "todo",
        "dep on": "",
        "name": "Simplify/rename env adapter API",
        "notes / blockers": "Reflect non-Gym-physics behavior in naming and surface.",
    },
    {
        "uid": "STR-003",
        "Sprint": "",
        "prio": "P3",
        "Size": "S",
        "status": "todo",
        "dep on": "",
        "name": "Remove deprecated run_simulation compatibility args",
        "notes / blockers": "After call sites migrated.",
    },
    {
        "uid": "STR-004",
        "Sprint": "S01",
        "prio": "P1",
        "Size": "M",
        "status": "wip",
        "dep on": "",
        "name": "Commit and split uncommitted WIP",
        "notes / blockers": "~3.5k LOC on feat/backend-stuff: notebooks/utils vs sim/render vs experiments vs cursor skills.",
    },
    {
        "uid": "FEAT-001",
        "Sprint": "S01",
        "prio": "P1",
        "Size": "L",
        "status": "wip",
        "dep on": "",
        "name": "Optimize simulation performance for fast ML iteration",
        "notes / blockers": "C/D ray batch promoted; sim_timing in place; export still bottleneck (~82s vs ~10s sim).",
    },
    {
        "uid": "FEAT-002",
        "Sprint": "S01",
        "prio": "P1",
        "Size": "S",
        "status": "todo",
        "dep on": "FEAT-001",
        "name": "Promote video export frame_stride=2",
        "notes / blockers": "Experiment supported ~2x export speedup; not yet in production render path.",
    },
    {
        "uid": "FEAT-003",
        "Sprint": "",
        "prio": "P3",
        "Size": "S",
        "status": "todo",
        "dep on": "",
        "name": "Telemetry: environment avg cloud speed",
        "notes / blockers": "",
    },
    {
        "uid": "FEAT-004",
        "Sprint": "S01",
        "prio": "P0",
        "Size": "M",
        "status": "wip",
        "dep on": "BUG-003",
        "name": "Complete S01 notebook 03 - movement constraints",
        "notes / blockers": "AttitudeSafetyController + unit tests green; e2e red; MP4 in artifacts/video_archive/. Plan: movement_constraints.plan.md",
    },
    {
        "uid": "FEAT-005",
        "Sprint": "S01",
        "prio": "P0",
        "Size": "M",
        "status": "wip",
        "dep on": "BUG-003",
        "name": "Complete S01 notebook 02 - clouds",
        "notes / blockers": "cloud_formation + verification utils; sim integration test red; visible motion TBD.",
    },
    {
        "uid": "FEAT-006",
        "Sprint": "S01",
        "prio": "P1",
        "Size": "XL",
        "status": "todo",
        "dep on": "FEAT-005",
        "name": "clouds_and_target_grid plan (production)",
        "notes / blockers": "Min observation time, stillness, area score, off-nadir damage, secondary FOV — .cursor/plans/v2/clouds_and_target_grid_808c8677.plan.md",
    },
    {
        "uid": "DONE-001",
        "Sprint": "S01",
        "prio": "P2",
        "Size": "M",
        "status": "done",
        "dep on": "",
        "name": "Sensor kernel: fused multi-camera ray batch (hypothesis C)",
        "notes / blockers": "Promoted 2026-06-03; ~1.6x micro, ~1.18x episode.",
    },
    {
        "uid": "DONE-002",
        "Sprint": "S01",
        "prio": "P2",
        "Size": "M",
        "status": "done",
        "dep on": "",
        "name": "camera_2d: vectorized tensor cloud x ray hits (hypothesis D)",
        "notes / blockers": "Promoted 2026-06-03; ~2.9x line kernel @ 63 clouds.",
    },
    {
        "uid": "DONE-003",
        "Sprint": "S01",
        "prio": "P2",
        "Size": "S",
        "status": "done",
        "dep on": "",
        "name": "sim_timing experiment harness",
        "notes / blockers": "High-cloud profile baseline; backend/scripts/experiments/sim_timing/",
    },
    {
        "uid": "DONE-004",
        "Sprint": "S01",
        "prio": "P2",
        "Size": "M",
        "status": "done",
        "dep on": "",
        "name": "Attitude safety controller + ATTITUDE_SAFETY constants",
        "notes / blockers": "Notebook 03 core logic; unit tests green.",
    },
)


def _column_index(sheet: Worksheet) -> dict[str, int]:
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    if header != list(BACKLOG_COLUMNS):
        raise ValueError(
            f"Backlog sheet header mismatch.\nExpected: {list(BACKLOG_COLUMNS)}\nGot:      {header}"
        )
    return {name: idx for idx, name in enumerate(BACKLOG_COLUMNS)}


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_backlog_entries(path: Path | str) -> list[dict[str, str]]:
    """Return backlog rows as dicts keyed by BACKLOG_COLUMNS."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if BACKLOG_SHEET not in workbook.sheetnames:
            raise ValueError(f"Missing sheet {BACKLOG_SHEET!r} in {path}")
        sheet = workbook[BACKLOG_SHEET]
        col_idx = _column_index(sheet)
        entries: list[dict[str, str]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            entry = {
                col: _cell_str(row[col_idx[col]] if col_idx[col] < len(row) else None)
                for col in BACKLOG_COLUMNS
            }
            entries.append(entry)
        return entries
    finally:
        workbook.close()


def _validate_entry(entry: dict[str, str], *, row_number: int) -> None:
    uid = entry["uid"]
    if not uid:
        raise ValueError(f"Row {row_number}: uid is required")
    status = entry["status"]
    if status and status not in ALLOWED_STATUS:
        raise ValueError(f"Row {row_number} ({uid}): invalid status {status!r}")
    prio = entry["prio"]
    if prio and prio not in ALLOWED_PRIO:
        raise ValueError(f"Row {row_number} ({uid}): invalid prio {prio!r}")
    size = entry["Size"]
    if size and size not in ALLOWED_SIZE:
        raise ValueError(f"Row {row_number} ({uid}): invalid Size {size!r}")


def assert_backlog_workbook_healthy(path: Path | str) -> None:
    """Raise if the workbook is missing, malformed, or has invalid rows."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Backlog workbook not found: {path}")

    entries = read_backlog_entries(path)
    if not entries:
        raise ValueError(f"Backlog workbook has no data rows: {path}")

    seen_uids: set[str] = set()
    for row_number, entry in enumerate(entries, start=2):
        _validate_entry(entry, row_number=row_number)
        uid = entry["uid"]
        if uid in seen_uids:
            raise ValueError(f"Duplicate uid {uid!r}")
        seen_uids.add(uid)


def write_backlog_workbook(
    path: Path | str,
    rows: tuple[dict[str, str], ...] | list[dict[str, str]],
    *,
    overwrite: bool = False,
) -> Path:
    """Write the backlog sheet (creates parent dirs if needed)."""
    path = Path(path)
    if path.exists() and not overwrite:
        raise FileExistsError(f"Backlog workbook already exists: {path} (use overwrite=True)")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = BACKLOG_SHEET
    sheet.append(list(BACKLOG_COLUMNS))
    for row in rows:
        sheet.append([row.get(col, "") for col in BACKLOG_COLUMNS])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def init_default_backlog(path: Path | str, *, overwrite: bool = False) -> Path:
    """Create backlog.xlsx with seeded rows from the current project state."""
    path = Path(path)
    write_backlog_workbook(path, DEFAULT_BACKLOG_ROWS, overwrite=overwrite)
    assert_backlog_workbook_healthy(path)
    return path


def _print_entries(entries: list[dict[str, str]]) -> None:
    active = [e for e in entries if e["status"] in {"todo", "wip", "blocked"}]
    done = [e for e in entries if e["status"] == "done"]
    print(f"entries={len(entries)} active={len(active)} done={len(done)}")
    for entry in entries:
        print(
            f"{entry['uid']:8}  {entry['status']:7}  {entry['prio']:2}  {entry['Sprint']:4}  {entry['name']}"
        )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Project backlog workbook utilities.")
    parser.add_argument(
        "command",
        choices=("check", "init", "list"),
        help="check=validate workbook; init=create default; list=print rows",
    )
    parser.add_argument(
        "--path",
        type=Path,
        default=None,
        help="Workbook path (default: ENV.PATHS.BACKLOG_XLSX)",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Overwrite existing workbook (init only)",
    )
    args = parser.parse_args(argv)

    if args.path is None:
        from ENV.PATHS import BACKLOG_XLSX

        path = BACKLOG_XLSX
    else:
        path = args.path

    if args.command == "init":
        init_default_backlog(path, overwrite=args.force)
        print(f"Created {path}")
        return 0

    if args.command == "check":
        assert_backlog_workbook_healthy(path)
        entries = read_backlog_entries(path)
        print(f"OK: {path} ({len(entries)} rows)")
        return 0

    entries = read_backlog_entries(path)
    _print_entries(entries)
    return 0


if __name__ == "__main__":
    sys.exit(main())
